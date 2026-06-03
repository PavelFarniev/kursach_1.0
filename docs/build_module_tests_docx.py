from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = Path("/Users/artemfarniev2013/WebstormProjects/kursach_1.0/docs/tests_killexam_modules.docx")

PROJECT_META = {
    "Название проекта": "KillExam",
    "Рабочая версия": "1.0",
    "Имя тестирующего": "Артем Фарниев",
    "Дата теста": "29.05.2026",
}


FIELD_DESCRIPTIONS = [
    ("Название проекта", "Название тестируемого проекта"),
    ("Рабочая версия", "Версия проекта / программного обеспечения"),
    ("Имя тестирующего", "Имя того, кто проводил тесты"),
    ("Дата(ы) теста", "Дата или набор дат проведения тестов"),
    ("Тестовый пример #", "Уникальный ID тестового примера. Например, TC_THEME_01 означает первый тест модуля темы."),
    ("Приоритет тестирования (Низкий/Средний/Высокий)", "Насколько важен тест для бизнес-логики и пользовательского сценария"),
    ("Заголовок/название теста", "Краткое название тестового случая"),
    ("Краткое изложение теста", "Описание того, что проверяется этим тестом"),
    ("Этапы теста", "Подробная последовательность действий пользователя"),
    ("Тестовые данные", "Данные, которые используются в тесте"),
    ("Ожидаемый результат", "Каким должен быть результат при корректной работе системы"),
    ("Фактический результат", "Что получилось после реального выполнения теста"),
    ("Предварительное условие", "Что должно быть выполнено до начала теста"),
    ("Постусловие", "Состояние системы после завершения теста"),
    ("Статус (Зачет/Незачет)", "Соответствует ли фактический результат ожидаемому"),
    ("Примечания/комментарии", "Дополнительные замечания по особенностям теста"),
]


@dataclass
class TestCase:
    test_id: str
    priority: str
    title: str
    summary: str
    steps: list[str]
    test_data: str
    expected: str
    actual: str
    status: str
    precondition: str
    postcondition: str
    notes: str


TEST_CASES = [
    TestCase(
        test_id="TC_THEME_01",
        priority="Высокий",
        title="Переключение светлой и тёмной темы на сайте",
        summary="Проверка, что пользователь может переключать тему интерфейса и оформление страницы изменяется без перезагрузки.",
        steps=[
            "Открыть главную страницу платформы.",
            "Найти кнопку переключения темы в верхней панели.",
            "Нажать на кнопку переключения темы.",
            "Проверить изменение фона, цвета текста, кнопок и карточек.",
            "Повторно нажать на кнопку и убедиться, что тема возвращается в исходное состояние.",
        ],
        test_data="Главная страница сайта KillExam.",
        expected="Тема интерфейса переключается между светлой и тёмной, оформление обновляется сразу, все элементы остаются читаемыми.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Открыта главная страница платформы.",
        postcondition="На сайте установлена последняя выбранная пользователем тема.",
        notes="Особое внимание следует уделить читаемости текста и видимости иконок в обеих темах.",
    ),
    TestCase(
        test_id="TC_THEME_02",
        priority="Высокий",
        title="Сохранение выбранной темы после обновления страницы",
        summary="Проверка, что выбранная пользователем тема сохраняется после повторной загрузки сайта.",
        steps=[
            "Открыть сайт и переключить его в тёмную тему.",
            "Обновить страницу браузера.",
            "Проверить оформление интерфейса после повторной загрузки.",
        ],
        test_data="Выбранная тема: тёмная.",
        expected="После обновления страницы сохраняется ранее выбранная пользователем тема.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь уже переключил сайт в одну из тем.",
        postcondition="Выбранная тема остаётся сохранённой в пользовательских настройках браузера.",
        notes="Аналогично следует проверить и сохранение светлой темы.",
    ),
    TestCase(
        test_id="TC_THEME_03",
        priority="Средний",
        title="Корректное отображение страницы входа в светлой теме",
        summary="Проверка, что при светлой теме страница входа отображается корректно и все элементы формы остаются доступными.",
        steps=[
            "Открыть страницу входа.",
            "Убедиться, что выбрана светлая тема.",
            "Проверить заголовок, поля Email и Пароль, кнопку входа и текстовые подписи.",
        ],
        test_data="Страница входа `/login`.",
        expected="Форма входа отображается без визуальных ошибок, текст читается, поля и кнопка доступны пользователю.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Frontend сайта запущен, страница входа доступна.",
        postcondition="Форма входа остаётся доступной для ввода данных.",
        notes="При необходимости дополнительно проверить страницу входа на узком экране.",
    ),
    TestCase(
        test_id="TC_PULSE_01",
        priority="Высокий",
        title="Отображение панели Pulse после авторизации",
        summary="Проверка, что после входа в систему пользователю становится доступна персональная панель статистики Pulse.",
        steps=[
            "Открыть страницу входа.",
            "Ввести корректные учётные данные пользователя.",
            "Нажать кнопку «Войти».",
            "Перейти на главную страницу системы.",
            "Проверить наличие блока Pulse.",
        ],
        test_data="Email: demo@student.ai\nПароль: demo123",
        expected="После авторизации на главной странице отображается персональная панель Pulse.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь не авторизован, страница входа доступна.",
        postcondition="Пользователь вошёл в систему, панель Pulse доступна на главной странице.",
        notes="Если авторизация не выполнена, вместо панели должна отображаться информационная заглушка.",
    ),
    TestCase(
        test_id="TC_PULSE_02",
        priority="Высокий",
        title="Отображение основных показателей Pulse",
        summary="Проверка, что в панели Pulse отображаются ключевые показатели статистики подготовки.",
        steps=[
            "Авторизоваться в системе.",
            "Открыть главную страницу.",
            "Найти панель Pulse.",
            "Проверить наличие показателей: готовность, активность, входы, серия.",
            "Проверить наличие блока динамики и рекомендаций.",
        ],
        test_data="Аккаунт пользователя с доступной панелью Pulse.",
        expected="В панели Pulse отображаются показатели готовности, активности, входов, серии, блок динамики и рекомендации.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь авторизован.",
        postcondition="Статистическая панель остаётся доступной до выхода из аккаунта.",
        notes="Допускается изменение числовых значений, если пользовательская активность изменилась ранее.",
    ),
    TestCase(
        test_id="TC_PULSE_03",
        priority="Высокий",
        title="Обновление статистики Pulse после учебной активности",
        summary="Проверка, что после прохождения слайдов курса показатели Pulse обновляются в соответствии с действиями пользователя.",
        steps=[
            "Авторизоваться в системе.",
            "Открыть курс «ЕГЭ Математика: профиль».",
            "Перейти в режим обучения.",
            "Пройти несколько слайдов курса.",
            "Вернуться на главную страницу и сравнить показатели Pulse.",
        ],
        test_data="Курс: ЕГЭ Математика: профиль",
        expected="После учебной активности изменяются показатели, связанные с прогрессом или активностью пользователя.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь авторизован и имеет доступ к курсу.",
        postcondition="Новая активность учтена в панели Pulse.",
        notes="Изменение конкретных значений зависит от накопленного состояния аккаунта.",
    ),
    TestCase(
        test_id="TC_PULSE_04",
        priority="Средний",
        title="Отображение рекомендаций и слабой зоны в Pulse",
        summary="Проверка, что в панели Pulse выводятся рекомендации по дальнейшему обучению и слабая зона пользователя.",
        steps=[
            "Авторизоваться в системе.",
            "Открыть главную страницу.",
            "Найти блок Pulse.",
            "Проверить наличие текста слабой зоны.",
            "Проверить наличие рекомендаций по дальнейшим действиям.",
        ],
        test_data="Панель Pulse авторизованного пользователя.",
        expected="В панели Pulse отображаются слабая зона и рекомендации, связанные с текущим прогрессом пользователя.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь авторизован, Pulse загружен.",
        postcondition="Рекомендации остаются доступными до изменения учебной активности.",
        notes="Если заметок пользователя нет, это не должно ломать отображение остальной панели.",
    ),
    TestCase(
        test_id="TC_FINAL_01",
        priority="Высокий",
        title="Переход к итоговому тестированию на последнем слайде курса",
        summary="Проверка, что после прохождения учебных материалов пользователь попадает на итоговый тест курса.",
        steps=[
            "Авторизоваться в системе.",
            "Открыть курс «ЕГЭ Математика: профиль».",
            "Перейти в режим обучения.",
            "Последовательно пролистать все учебные слайды.",
            "Проверить последний экран курса.",
        ],
        test_data="Курс: ЕГЭ Математика: профиль",
        expected="На последнем этапе курса отображается экран итогового тестирования.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь авторизован и открыл курс.",
        postcondition="Пользователь находится на экране итогового теста.",
        notes="На последнем слайде должен отображаться незавершённый статус курса до проверки ответов.",
    ),
    TestCase(
        test_id="TC_FINAL_02",
        priority="Высокий",
        title="Блокировка завершения курса до прохождения итогового теста",
        summary="Проверка, что курс не считается завершённым, пока пользователь не прошёл итоговый тест.",
        steps=[
            "Открыть последний экран курса с итоговым тестом.",
            "Не отправляя ответы, проверить текущий прогресс курса.",
            "Проверить верхний блок статуса курса.",
        ],
        test_data="Последний экран курса до отправки теста.",
        expected="Прогресс курса отображается как 99%, курс не имеет статуса завершённого.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь дошёл до финального теста курса.",
        postcondition="Курс остаётся в состоянии ожидания прохождения итогового теста.",
        notes="Только успешное прохождение теста должно переводить прогресс на 100%.",
    ),
    TestCase(
        test_id="TC_FINAL_03",
        priority="Высокий",
        title="Неуспешное прохождение итогового теста",
        summary="Проверка реакции системы на неправильные ответы пользователя в итоговом тесте.",
        steps=[
            "Открыть экран итогового тестирования.",
            "Выбрать заведомо неправильные ответы.",
            "Нажать кнопку «Проверить ответы».",
            "Проверить сообщение системы и статус курса.",
        ],
        test_data="Неверные варианты ответов для всех вопросов теста.",
        expected="Система сообщает, что тест не пройден, показывает результат и предлагает повторную попытку, при этом курс не завершается.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь находится на финальном тесте курса.",
        postcondition="Курс остаётся незавершённым, доступна кнопка повторного прохождения.",
        notes="Ошибка пользователя не должна ломать расположение элементов интерфейса.",
    ),
    TestCase(
        test_id="TC_FINAL_04",
        priority="Средний",
        title="Повторное прохождение итогового теста после неуспешной попытки",
        summary="Проверка, что после неуспешной попытки пользователь может пройти тест заново.",
        steps=[
            "Получить отрицательный результат итогового теста.",
            "Нажать кнопку повторного прохождения теста.",
            "Убедиться, что ответы можно выбрать заново.",
            "Повторно пройти тест.",
        ],
        test_data="Неуспешно завершённый итоговый тест.",
        expected="Тест сбрасывается, пользователь получает возможность выбрать ответы повторно.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь уже выполнил одну неуспешную попытку тестирования.",
        postcondition="Форма теста снова доступна для заполнения.",
        notes="При повторном открытии не должны оставаться заблокированные ответы от предыдущей попытки.",
    ),
    TestCase(
        test_id="TC_FINAL_05",
        priority="Высокий",
        title="Успешное завершение курса после итогового тестирования",
        summary="Проверка, что при правильных ответах итоговый тест переводит курс в завершённое состояние.",
        steps=[
            "Открыть итоговый тест курса.",
            "Выбрать правильные ответы на все вопросы.",
            "Нажать кнопку «Проверить ответы».",
            "Проверить итоговое сообщение и верхний блок прогресса.",
        ],
        test_data="Правильные ответы на итоговый тест по курсу «ЕГЭ Математика: профиль».",
        expected="Система сообщает об успешном прохождении теста, прогресс курса становится равным 100%, курс получает статус завершённого.",
        actual="Соответствует ожидаемому результату.",
        status="Зачет",
        precondition="Пользователь находится на экране итогового тестирования.",
        postcondition="Курс завершён, финальный тест можно открыть повторно без сброса зачёта.",
        notes="После успешного прохождения должна быть доступна самопроверка без потери статуса завершения.",
    ),
]


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))


def style_paragraph(paragraph, *, bold: bool = False, italic: bool = False, size: int = 14, color: str | None = None, align=WD_ALIGN_PARAGRAPH.LEFT, space_after: int = 0) -> None:
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text: str, *, bold: bool = False, size: int = 14) -> None:
    cell.text = ""
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        style_paragraph(paragraph, bold=bold, size=size)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.bold = bold
        run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def build_info_table(document: Document) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    for key, value in PROJECT_META.items():
        row = table.add_row().cells
        set_cell_width(row[0], 5.5)
        set_cell_width(row[1], 5.5)
        set_cell_text(row[0], key)
        set_cell_text(row[1], value)

    set_table_borders(table)


def build_field_description_table(document: Document) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    header = table.rows[0].cells
    set_cell_width(header[0], 8.0)
    set_cell_width(header[1], 8.0)
    set_cell_text(header[0], "Поле", bold=True)
    set_cell_text(header[1], "Описание", bold=True)

    for key, value in FIELD_DESCRIPTIONS:
        row = table.add_row().cells
        set_cell_width(row[0], 8.0)
        set_cell_width(row[1], 8.0)
        set_cell_text(row[0], key)
        set_cell_text(row[1], value)

    set_table_borders(table)


def add_test_case_table(document: Document, test_number: int, case: TestCase) -> None:
    label = document.add_paragraph()
    label.add_run(f"Тестовый пример #{test_number}:")
    style_paragraph(label, italic=True, size=14, space_after=3)

    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    rows = [
        ("Тестовый пример #", case.test_id),
        ("Приоритет тестирования", case.priority),
        ("Заголовок/название теста", case.title),
        ("Краткое изложение теста", case.summary),
        ("Этапы теста", "\n".join(f"{index}. {step}" for index, step in enumerate(case.steps, start=1))),
        ("Тестовые данные", case.test_data),
        ("Ожидаемый результат", case.expected),
        ("Фактический результат", case.actual),
        ("Статус", case.status),
        ("Предварительное условие", case.precondition),
        ("Постусловие", case.postcondition),
        ("Примечания/комментарии", case.notes),
    ]

    for key, value in rows:
        row = table.add_row().cells
        set_cell_width(row[0], 7.2)
        set_cell_width(row[1], 11.2)
        set_cell_text(row[0], key)
        set_cell_text(row[1], value)

    set_table_borders(table)
    document.add_paragraph()


def build_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(14)

    title = document.add_paragraph()
    title.add_run("Тесты")
    style_paragraph(title, bold=True, size=18, space_after=8)

    heading = document.add_paragraph()
    heading.add_run("Аннотация теста")
    style_paragraph(heading, bold=True, size=16, color="2F5597", space_after=6)
    build_info_table(document)

    document.add_paragraph()

    heading2 = document.add_paragraph()
    heading2.add_run("Расшифровка тестовых информационных полей")
    style_paragraph(heading2, bold=True, size=16, color="2F5597", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    build_field_description_table(document)

    document.add_paragraph()

    for index, case in enumerate(TEST_CASES, start=1):
        add_test_case_table(document, index, case)

    return document


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
