from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
SOURCE_PATH = BASE_DIR / "отчёт.txt"
OUTPUT_PATH = BASE_DIR / "отчёт_AI_модуль_ГОСТ_2017.docx"


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def configure_document(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "FigureCaption" not in styles:
        style = styles.add_style("FigureCaption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "ContentsLine" not in styles:
        style = styles.add_style("ContentsLine", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def add_paragraph(
    document,
    text="",
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line=True,
    bold=False,
    size=14,
    style=None,
):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(1.25) if first_line else Cm(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    for run in paragraph.runs:
        set_run_font(run, bold=True)
    return paragraph


def add_centered(document, text, *, bold=False, size=14):
    return add_paragraph(
        document,
        text,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line=False,
        bold=bold,
        size=size,
    )


def is_main_heading(line):
    return (
        line
        in {
            "ВВЕДЕНИЕ",
            "ЗАКЛЮЧЕНИЕ",
            "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            "ПРИЛОЖЕНИЕ А",
        }
        or line.startswith("1 Теоретическая часть")
        or line.startswith("2 Технологическая часть")
        or line.startswith("3 Практическая часть")
    )


def is_sub_heading(line):
    return (
        line.startswith("1.1 ")
        or line.startswith("1.2 ")
        or line.startswith("1.3 ")
        or line.startswith("2.1 ")
        or line.startswith("2.2 ")
    )


def is_figure_caption(line):
    return line.startswith("Рисунок ")


def add_title_page(document):
    add_centered(document, "ОТЧЕТ", bold=True, size=16)
    add_centered(
        document,
        "по разработке программного продукта",
        bold=True,
        size=14,
    )
    add_centered(
        document,
        "Тема: разработка AI-модуля для платформы подготовки к экзаменам",
        bold=True,
        size=14,
    )
    add_paragraph(document, "", first_line=False)
    add_paragraph(document, "", first_line=False)
    add_centered(document, "AI-модуль платформы подготовки к экзаменам", size=14)
    document.add_page_break()


def build_document():
    source_text = SOURCE_PATH.read_text(encoding="utf-8")
    lines = [line.rstrip() for line in source_text.splitlines()]

    document = Document()
    configure_document(document)
    add_page_number(document.sections[0])
    add_title_page(document)

    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.strip()

        if not line:
            i += 1
            continue

        if line == "СОДЕРЖАНИЕ":
            add_heading(document, "СОДЕРЖАНИЕ", level=1)
            i += 1
            while i < len(lines) and lines[i].strip() != "ВВЕДЕНИЕ":
                content_line = lines[i].strip()
                if content_line:
                    add_paragraph(
                        document, content_line, first_line=False, style="ContentsLine"
                    )
                i += 1
            document.add_page_break()
            continue

        if is_main_heading(line):
            add_heading(document, line, level=1)
            i += 1
            continue

        if is_sub_heading(line):
            add_heading(document, line, level=2)
            i += 1
            continue

        if is_figure_caption(line):
            add_centered(document, "[Место для вставки диаграммы]", bold=True)
            add_paragraph(document, line, first_line=False, style="FigureCaption")
            i += 1
            continue

        if line.startswith(tuple(f"{number}." for number in range(1, 20))):
            add_paragraph(document, line, first_line=False)
            i += 1
            continue

        add_paragraph(document, line)
        i += 1

    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
